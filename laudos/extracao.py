"""Leitura do documento de anamnese enviado pela psicóloga: PDF, Word ou texto."""

import io
from pathlib import Path


class ArquivoInvalido(Exception):
    pass


def texto_do_arquivo(arquivo):
    nome = arquivo.name.lower()
    dados = arquivo.read()
    sufixo = Path(nome).suffix
    if sufixo == ".pdf":
        from pypdf import PdfReader

        leitor = PdfReader(io.BytesIO(dados))
        texto = "\n".join((pagina.extract_text() or "") for pagina in leitor.pages)
    elif sufixo == ".docx":
        import docx

        documento = docx.Document(io.BytesIO(dados))
        linhas = [p.text for p in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                linhas.append(" | ".join(celula.text for celula in linha.cells))
        texto = "\n".join(linhas)
    elif sufixo in (".txt", ".md"):
        texto = dados.decode("utf-8", errors="replace")
    else:
        raise ArquivoInvalido("Envie a anamnese em PDF, Word (.docx) ou texto (.txt).")
    texto = texto.strip()
    if not texto:
        raise ArquivoInvalido(
            "Não consegui ler texto nesse arquivo. Se for um PDF escaneado, cole o texto da anamnese no campo abaixo."
        )
    return texto
